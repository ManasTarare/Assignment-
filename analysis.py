"""
Trader Performance vs Bitcoin Market Sentiment - Analysis Pipeline
Inputs:
  - fear_greed_index.csv  (columns: timestamp, value, classification, date)
  - historical_data.csv   (Hyperliquid trade data)
Outputs:
  - merged.csv
  - chart1_avg_pnl.png, chart2_winrate.png, chart3_size.png,
    chart4_totalpnl.png, chart5_longshort.png
  - trader_sentiment_analysis.xlsx
"""

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FG_PATH = 'fear_greed_index.csv'
HD_PATH = 'historical_data.csv'

# ---------------------------------------------------------------------
# 1. Load and merge
# ---------------------------------------------------------------------
fg = pd.read_csv(FG_PATH)
hd = pd.read_csv(HD_PATH)

hd['date'] = pd.to_datetime(hd['Timestamp IST'], format='%d-%m-%Y %H:%M').dt.date
fg['date'] = pd.to_datetime(fg['date']).dt.date

merged = hd.merge(fg[['date', 'classification']], on='date', how='left')
print("Unmatched rows:", merged['classification'].isna().sum(), "of", len(merged))
merged.to_csv('merged.csv', index=False)

# ---------------------------------------------------------------------
# 2. Derived columns
# ---------------------------------------------------------------------
merged = merged.dropna(subset=['classification'])

order = ['Extreme Fear', 'Fear', 'Neutral', 'Greed', 'Extreme Greed']
merged['classification'] = pd.Categorical(merged['classification'], categories=order, ordered=True)

def broad(c):
    c = str(c)
    if 'Fear' in c:
        return 'Fear'
    if 'Greed' in c:
        return 'Greed'
    return 'Neutral'

merged['sentiment_broad'] = merged['classification'].astype(str).apply(broad)

# "Closed trades" = rows with a realized PnL event
closed = merged[merged['Closed PnL'] != 0]

# ---------------------------------------------------------------------
# 3. Summary tables
# ---------------------------------------------------------------------
def profit_factor(x):
    profit = x[x > 0].sum()
    loss = -x[x < 0].sum()
    return profit / loss if loss > 0 else np.nan

g = closed.groupby('classification').agg(
    Trades=('Closed PnL', 'count'),
    Total_PnL=('Closed PnL', 'sum'),
    Avg_PnL=('Closed PnL', 'mean'),
    Median_PnL=('Closed PnL', 'median'),
    Win_Rate=('Closed PnL', lambda x: (x > 0).mean()),
    Avg_Size_USD=('Size USD', 'mean'),
).reindex(order).round(3)

g2 = closed.groupby('sentiment_broad').agg(
    Trades=('Closed PnL', 'count'),
    Total_PnL=('Closed PnL', 'sum'),
    Avg_PnL=('Closed PnL', 'mean'),
    Win_Rate=('Closed PnL', lambda x: (x > 0).mean()),
    Avg_Size_USD=('Size USD', 'mean'),
).reindex(['Fear', 'Neutral', 'Greed']).round(3)

pf_class = closed.groupby('classification')['Closed PnL'].apply(profit_factor).reindex(order).round(3)

print("\nBy classification:\n", g)
print("\nBy broad sentiment:\n", g2)
print("\nProfit factor by classification:\n", pf_class)

# Long/Short opening behavior by broad sentiment
opens = merged[merged['Direction'].isin(['Open Long', 'Open Short'])]
ls_ct = pd.crosstab(opens['sentiment_broad'], opens['Direction'], normalize='index') * 100
ls_ct = ls_ct.reindex(['Fear', 'Neutral', 'Greed'])
print("\nLong/Short open % by sentiment:\n", ls_ct)

# ---------------------------------------------------------------------
# 4. Charts
# ---------------------------------------------------------------------
colors = ['#8B0000', '#E07B39', '#9E9E9E', '#6FA85A', '#1B5E20']

def bar_chart(series, title, ylabel, fname, ylim=None):
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.bar(series.index.astype(str), series.values, color=colors[:len(series)])
    ax.set_title(title)
    ax.set_ylabel(ylabel)
    if ylim:
        ax.set_ylim(*ylim)
    plt.xticks(rotation=20)
    plt.tight_layout()
    plt.savefig(fname, dpi=150)
    plt.close()

bar_chart(g['Avg_PnL'], 'Average Closed PnL per Trade by Market Sentiment',
          'Avg Closed PnL (USD)', 'chart1_avg_pnl.png')

bar_chart(g['Win_Rate'] * 100, 'Win Rate (%) by Market Sentiment',
          'Win Rate (%)', 'chart2_winrate.png', ylim=(0, 100))

bar_chart(g['Avg_Size_USD'], 'Average Trade Size (USD) by Market Sentiment',
          'Avg Size (USD)', 'chart3_size.png')

bar_chart(g['Total_PnL'], 'Total Closed PnL (USD) by Market Sentiment',
          'Total Closed PnL (USD)', 'chart4_totalpnl.png')

# Long vs short stacked bar
fig, ax = plt.subplots(figsize=(7, 4.5))
ls_ct.plot(kind='bar', stacked=True, ax=ax, color=['#1B5E20', '#8B0000'])
ax.set_title('Long vs Short Position Openings by Sentiment')
ax.set_ylabel('% of Opening Trades')
plt.xticks(rotation=0)
plt.legend(title='')
plt.tight_layout()
plt.savefig('chart5_longshort.png', dpi=150)
plt.close()

# ---------------------------------------------------------------------
# 5. Excel workbook
# ---------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Summary by Classification"

header_font = Font(bold=True, color="FFFFFF", name="Arial")
header_fill = PatternFill("solid", start_color="1B5E20")
title_font = Font(bold=True, size=14, name="Arial")
border = Border(*[Side(style='thin', color='CCCCCC')] * 4)

ws['A1'] = "Trader Performance by Market Sentiment (Fear & Greed Index)"
ws['A1'].font = title_font
ws.append([])
ws.append(['Classification'] + list(g.columns) + ['Profit Factor'])
for c in ws[3]:
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')

for cls in order:
    ws.append([cls] + list(g.loc[cls]) + [pf_class.loc[cls]])

for r in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for c in r:
        c.border = border
        if c.column > 1:
            c.number_format = '#,##0.00'

for col, w in zip('ABCDEFGH', [16, 10, 16, 12, 12, 11, 14, 12]):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet("Summary by Broad Sentiment")
ws2['A1'] = "Trader Performance: Fear vs Neutral vs Greed"
ws2['A1'].font = title_font
ws2.append([])
ws2.append(['Sentiment'] + list(g2.columns))
for c in ws2[3]:
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')

for s in ['Fear', 'Neutral', 'Greed']:
    ws2.append([s] + list(g2.loc[s]))

for r in ws2.iter_rows(min_row=3, max_row=ws2.max_row, min_col=1, max_col=ws2.max_column):
    for c in r:
        c.border = border
        if c.column > 1:
            c.number_format = '#,##0.00'

for col, w in zip('ABCDEF', [14, 10, 16, 12, 12, 14]):
    ws2.column_dimensions[col].width = w

ws3 = wb.create_sheet("Notes")
ws3['A1'] = "Notes"
ws3['A1'].font = title_font
notes = [
    "Data: trades from Hyperliquid merged with Bitcoin Fear & Greed Index by date.",
    "'Closed trades' = rows where Closed PnL != 0 (realized PnL events), used for PnL/win-rate analysis.",
    "Profit Factor = Gross Profit / Gross Loss.",
    "Broad Sentiment groups: Extreme Fear+Fear -> Fear; Extreme Greed+Greed -> Greed; Neutral -> Neutral.",
]
for i, n in enumerate(notes, start=3):
    ws3.cell(row=i, column=1, value=n)
ws3.column_dimensions['A'].width = 110

wb.save('trader_sentiment_analysis.xlsx')
print("\nSaved trader_sentiment_analysis.xlsx")
